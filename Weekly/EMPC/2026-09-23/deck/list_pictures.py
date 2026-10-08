# -*- coding: utf-8 -*-
"""List every picture in the deck with its placed geometry, to sanity-check that
formula images are typeset at a sensible size rather than stretched to column width.
"""

import os
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
   HERE, "Capstone_Weekly_2026-09-23.pptx")

prs = Presentation(DECK)
print("%-4s %-6s %8s %8s %8s %8s  %s" %
      ("p", "shape", "left", "top", "width", "height", "native px / aspect"))
for idx, slide in enumerate(prs.slides, start=1):
   for shape in slide.shapes:
      if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
         continue
      img = shape.image
      w_in = shape.width / 914400
      h_in = shape.height / 914400
      print("%-4d %-6d %8.2f %8.2f %8.2f %8.2f  %dx%d  placed aspect %.3f" %
            (idx, shape.shape_id, shape.left / 914400, shape.top / 914400,
             w_in, h_in, img.size[0], img.size[1], w_in / h_in if h_in else 0))
